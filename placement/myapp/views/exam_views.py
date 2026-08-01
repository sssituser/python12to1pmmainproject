from rest_framework.decorators import api_view, permission_classes, throttle_classes, authentication_classes
from myapp.throttles import ExamRateThrottle, AuthenticatedUserThrottle
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework import status
from django.utils import timezone
from django.shortcuts import get_object_or_404
import json
import csv
import io
import re

from myapp.models import (
    ExamSession, ExamAnswer, PythonQuestion, Choice,
    AutomatedExamConfig, ExamQuestion, ExamQuestionChoice, ExamPaper,
    ExamPaperQuestionRelation, ExamViolationLog
)
from myapp.serializers import PythonQuestionSerializer, AutomatedExamConfigSerializer
from myapp.email_utils import send_exam_confirmation_email
import threading
import random


# ============================================================
# PLACEMENT EXAM DYNAMIC STORAGE
# ============================================================

_EXAM_STORE = []   # Simple in-process store (persists for session)

@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def create_placement_exam(request):
    """
    Persist a complete 10-step wizard exam to the database.
    """
    try:
        data = request.data

        title = str(data.get('title', '')).strip()
        if not title:
            return Response({"error": "Exam title is required"}, status=400)

        exam_type = data.get('exam_type', 'daily')
        subject = str(data.get('subject', 'PYTHON')).upper()
        course_name = str(data.get('course_name', '')).strip().upper()
        if not course_name:
            course_name = "ALL COURSES"

        def safe_int(val, default):
            try:
                if val is None or str(val).strip() == "": return default
                return int(float(val))
            except:
                return default

        duration = safe_int(data.get('duration'), 60)
        total_questions = safe_int(data.get('total_questions'), 30)
        total_marks = safe_int(data.get('total_marks'), 30)
        pass_marks = safe_int(data.get('pass_marks'), 15)
        marks_per_question = safe_int(data.get('marks_per_question'), 1)

        import json as _json
        unique_course_key = f"{course_name or subject}::{exam_type}::{title}"[:255]

        config_data = {
            'exam_name': title,
            'subjects': [subject],
            'duration': duration,
            'passing_strategy': 'marks',
            'requirement': pass_marks,
            'question_count': total_questions,
            'marks_per_question': marks_per_question,
        }

        config, created = AutomatedExamConfig.objects.update_or_create(
            course_name=unique_course_key,
            defaults=config_data
        )

        questions = data.get('questions', [])
        paper = None
        if questions:
            paper = ExamPaper.objects.create(
                title=title,
                subject=subject,
                duration=duration,
                total_marks=total_marks,
                instructions=data.get('description', '')
            )
            for idx, q in enumerate(questions):
                question = ExamQuestion.objects.create(
                    subject=subject,
                    topic=data.get('topic', ''),
                    difficulty=q.get('difficulty', 'medium'),
                    question_text=q.get('question', ''),
                    question_type='mcq',
                    marks=q.get('marks', marks_per_question)
                )
                opts = q.get('options', [])
                correct_idx = safe_int(q.get('correct'), 0)
                for oi, opt_text in enumerate(opts):
                    ExamQuestionChoice.objects.create(
                        question=question,
                        choice_text=opt_text,
                        is_correct=(oi == correct_idx)
                    )
                ExamPaperQuestionRelation.objects.create(
                    paper=paper, question=question, order=idx + 1
                )

        exam_entry = {
            'id': config.id,
            'title': title,
            'exam_type': exam_type,
            'subject': subject,
            'course': course_name,
            'duration': duration,
            'total_questions': len(questions) if questions else total_questions,
            'total_marks': total_marks,
            'pass_marks': pass_marks,
            'paper_id': paper.id if paper else None,
            'status': 'scheduled',
            'created_at': timezone.now().isoformat(),
            'settings': {
                'webcam_required': data.get('webcam_required', False),
                'face_detection': data.get('face_detection', True),
                'multi_face_detection': data.get('multi_face_detection', True),
                'fullscreen_required': data.get('fullscreen_required', True),
                'tab_switch_limit': data.get('tab_switch_limit', 3),
                'disable_copy_paste': data.get('disable_copy_paste', True),
                'randomize_questions': data.get('randomize_questions', True),
                'randomize_options': data.get('randomize_options', True),
                'negative_marking': data.get('negative_marking', False),
                'negative_marks': data.get('negative_marks', 0.25),
                'auto_submit': data.get('auto_submit', True),
                'risk_threshold': data.get('risk_threshold', 50),
                'departments': data.get('departments', []),
                'years': data.get('years', []),
                'start_time': data.get('start_time', ''),
                'end_time': data.get('end_time', ''),
                'show_result_immediately': data.get('show_result_immediately', True),
                'show_leaderboard': data.get('show_leaderboard', True),
                'certificate_enabled': data.get('certificate_enabled', False),
            }
        }
        global _EXAM_STORE
        _EXAM_STORE = [e for e in _EXAM_STORE if e['id'] != config.id]
        _EXAM_STORE.insert(0, exam_entry)

        # Sync to exam_settings.json
        try:
            import os
            base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            settings_file = os.path.join(base_dir, 'exam_settings.json')

            existing_settings = {}
            if os.path.exists(settings_file):
                with open(settings_file, 'r', encoding='utf-8') as sf:
                    content = sf.read().strip()
                    if content:
                        existing_settings = _json.loads(content)

            category = exam_type.capitalize()

            formatted_qs = []
            for idx, q in enumerate(questions):
                opts = q.get('options', [])
                correct_idx = safe_int(q.get('correct'), 0)
                answer_text = opts[correct_idx] if 0 <= correct_idx < len(opts) else ""

                formatted_qs.append({
                    "id": q.get('id') or int(timezone.now().timestamp() * 1000) + idx,
                    "question": q.get('question', ''),
                    "options": opts,
                    "answer": answer_text,
                    "type": q.get('type', 'mcq'),
                    "marks": safe_int(q.get('marks'), marks_per_question),
                    "subject": subject.upper()
                })

            exam_config = {
                "maxQuestions": len(formatted_qs) if formatted_qs else total_questions,
                "questions": formatted_qs,
                "passingRule": "percentage",
                "passingValue": int((pass_marks / total_marks) * 100) if total_marks > 0 else 50,
                "duration": duration
            }

            if course_name.upper() == "ALL COURSES":
                existing_settings[category] = exam_config
                for k in list(existing_settings.keys()):
                    if k.lower().endswith(f"_{category.lower()}"):
                        existing_settings[k] = exam_config
            else:
                target_key = f"{course_name}_{category}" if course_name else category
                resolved_key = target_key
                for k in existing_settings.keys():
                    if k.lower() == target_key.lower():
                        resolved_key = k
                        break
                existing_settings[resolved_key] = exam_config

            with open(settings_file, 'w', encoding='utf-8') as sf:
                _json.dump(existing_settings, sf, indent=4)
        except Exception as es_err:
            print(f"Error syncing to exam_settings.json: {es_err}")

        return Response({
            "status": "success",
            "exam_id": config.id,
            "paper_id": paper.id if paper else None,
            "message": f"Exam '{title}' published successfully!"
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({"status": "error", "message": f"Server error occurred: {str(e)}"}, status=500)


@api_view(['GET'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def list_placement_exams(request):
    """
    List all published placement exams (from in-memory store + DB)
    """
    global _EXAM_STORE

    db_configs = AutomatedExamConfig.objects.all().order_by('-id')[:50]
    db_entries = []
    for c in db_configs:
        parts = c.course_name.split("::")
        if len(parts) < 3:
            continue
        exam_type = parts[1]
        title = parts[2]
        paper = ExamPaper.objects.filter(title=title).order_by('-id').first()
        db_entries.append({
            'id': c.id,
            'title': title,
            'exam_type': exam_type,
            'subject': c.subjects[0] if c.subjects else 'N/A',
            'course': parts[0] if parts else c.course_name,
            'duration': c.duration,
            'total_questions': c.question_count,
            'total_marks': c.question_count * c.marks_per_question,
            'pass_marks': c.requirement,
            'paper_id': paper.id if paper else None,
            'start_time': None,
            'end_time': None,
            'status': 'scheduled',
            'settings': {
                'webcam_required': False,
                'face_detection': True,
                'fullscreen_required': True,
            }
        })

    mem_ids = {e['id'] for e in _EXAM_STORE}
    merged = list(_EXAM_STORE) + [e for e in db_entries if e['id'] not in mem_ids]

    return Response(merged)


@api_view(['GET'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def get_placement_exam_detail(request, exam_id):
    """
    Get detailed exam metadata and questions for student attempt.
    Also checks if student has already completed this exam.
    """
    global _EXAM_STORE

    user_identifier = ""
    if request.user and request.user.is_authenticated:
        user_identifier = request.user.email or request.user.username
    else:
        # Check authorization header token manually or student_email param
        auth_hdr = request.headers.get('Authorization', '')
        if auth_hdr.startswith('Bearer '):
            token_str = auth_hdr.split(' ')[1]
            try:
                from rest_framework_simplejwt.tokens import AccessToken
                from django.contrib.auth import get_user_model
                token = AccessToken(token_str)
                user_obj = get_user_model().objects.get(id=token['user_id'])
                user_identifier = user_obj.email or user_obj.username
            except Exception:
                pass
    
    if not user_identifier:
        user_identifier = request.query_params.get('student_email', '')

    # Check if student already has a completed session for this exam
    already_taken = False
    if user_identifier:
        session_exists = ExamSession.objects.filter(
            student_email=user_identifier,
            status='completed',
            student_name=str(exam_id)
        ).exists()
        if session_exists:
            already_taken = True

    # 1. Search in-memory store
    for e in _EXAM_STORE:
        if str(e['id']) == str(exam_id):
            exam_data = dict(e)
            exam_data['already_taken'] = already_taken
            if exam_data.get('paper_id'):
                try:
                    paper = ExamPaper.objects.get(id=exam_data['paper_id'])
                    rel_qs = ExamPaperQuestionRelation.objects.filter(paper=paper).order_by('order')
                    questions = []
                    for r in rel_qs:
                        q = r.question
                        choices = list(q.choices.values_list('choice_text', flat=True))
                        correct_idx = 0
                        for idx, c in enumerate(q.choices.all()):
                            if c.is_correct:
                                correct_idx = idx
                                break
                        questions.append({
                            'id': q.id,
                            'question_text': q.question_text,
                            'options': choices,
                            'correct_option_index': correct_idx,
                            'difficulty': q.difficulty,
                            'marks': q.marks,
                        })
                    exam_data['questions'] = questions
                except Exception as ex:
                    print(f"Error fetching questions for paper: {ex}")
            return Response(exam_data)

    # 2. Fallback to AutomatedExamConfig
    try:
        config = AutomatedExamConfig.objects.get(id=exam_id)
        parts = config.course_name.split("::")
        title = parts[2] if len(parts) >= 3 else config.exam_name
        paper = ExamPaper.objects.filter(title=title).order_by('-id').first()
        questions = []
        if paper:
            rel_qs = ExamPaperQuestionRelation.objects.filter(paper=paper).order_by('order')
            for r in rel_qs:
                q = r.question
                choices = list(q.choices.values_list('choice_text', flat=True))
                correct_idx = 0
                for idx, c in enumerate(q.choices.all()):
                    if c.is_correct:
                        correct_idx = idx
                        break
                questions.append({
                    'id': q.id,
                    'question_text': q.question_text,
                    'options': choices,
                    'correct_option_index': correct_idx,
                    'difficulty': q.difficulty,
                    'marks': q.marks,
                })

        exam_data = {
            'id': config.id,
            'title': title,
            'exam_type': parts[1] if len(parts) >= 3 else 'daily',
            'subject': config.subjects[0] if config.subjects else 'PYTHON',
            'duration': config.duration,
            'total_questions': len(questions) or config.question_count,
            'total_marks': config.question_count * config.marks_per_question,
            'pass_marks': config.requirement,
            'already_taken': already_taken,
            'questions': questions,
            'settings': {
                'webcam_required': False,
                'face_detection': True,
                'fullscreen_required': True,
                'disable_copy_paste': True,
                'disable_right_click': True,
                'tab_switch_limit': 3,
                'auto_submit': True
            }
        }
        return Response(exam_data)
    except AutomatedExamConfig.DoesNotExist:
        return Response({"detail": "Exam not found"}, status=404)


@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def submit_placement_exam(request, exam_id):
    """
    Save student's exam result submission and mark exam attempt as completed.
    """
    try:
        data = request.data
        user_identifier = ""
        if request.user and request.user.is_authenticated:
            user_identifier = request.user.email or request.user.username
        else:
            auth_hdr = request.headers.get('Authorization', '')
            if auth_hdr.startswith('Bearer '):
                token_str = auth_hdr.split(' ')[1]
                try:
                    from rest_framework_simplejwt.tokens import AccessToken
                    from django.contrib.auth import get_user_model
                    token = AccessToken(token_str)
                    user_obj = get_user_model().objects.get(id=token['user_id'])
                    user_identifier = user_obj.email or user_obj.username
                except Exception:
                    pass
        
        if not user_identifier:
            user_identifier = data.get('student_email', 'anonymous@student.com')

        score = data.get('score', 0)
        total = data.get('total', 100)
        percentage = round((score / total) * 100, 2) if total > 0 else 0
        
        # Record completed ExamSession for one-time enforcement
        ExamSession.objects.create(
            student_name=str(exam_id),
            student_email=user_identifier,
            start_time=timezone.now(),
            end_time=timezone.now(),
            status='completed',
            score=int(score),
            total_marks=int(total)
        )

        report_entry = {
            'exam_id': exam_id,
            'score': score,
            'total': total,
            'percentage': percentage,
            'passed': data.get('passed', False),
            'correct': data.get('correct', 0),
            'wrong': data.get('wrong', 0),
            'unattempted': data.get('unattempted', 0),
            'time_taken': data.get('time_taken', 0),
            'submitted_at': timezone.now().isoformat()
        }
        
        return Response({
            "status": "success",
            "message": "Exam result saved successfully!",
            "result": report_entry
        })
    except Exception as e:
        return Response({"error": str(e)}, status=500)


@api_view(['DELETE'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def delete_placement_exam(request, exam_id):
    """
    Delete a placement exam from both in-memory store and DB.
    """
    global _EXAM_STORE

    before = len(_EXAM_STORE)
    _EXAM_STORE = [e for e in _EXAM_STORE if e['id'] != exam_id]

    deleted_count, _ = AutomatedExamConfig.objects.filter(id=exam_id).delete()

    if deleted_count == 0 and len(_EXAM_STORE) == before:
        return Response({"error": "Exam not found"}, status=404)

    return Response({"status": "success", "message": "Exam deleted successfully"}, status=200)


# ---------------- MULTI-SUBJECT EXAM ENGINE APIS ----------------

@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([ExamRateThrottle])
def log_exam_violation(request):
    """
    Log secure assessment visibility, tab switches, and full-screen exits.
    """
    data = request.data
    session_id = data.get('session_id')
    violation_type = data.get('violation_type')
    remarks = data.get('remarks', '')
    image_snapshot = data.get('image_snapshot', '')

    if not session_id or not violation_type:
        return Response({"error": "session_id and violation_type are required"}, status=400)

    session = get_object_or_404(ExamSession, id=session_id)

    penalty_mapping = {
        'TAB_SWITCH': 5,
        'FULLSCREEN_EXIT': 10,
        'FACE_MISSING': 10,
        'MULTIPLE_FACE': 25,
        'PHONE_DETECTED': 30,
        'COPY_ATTEMPT': 5,
    }
    severity = penalty_mapping.get(violation_type, 5)

    violation = ExamViolationLog.objects.create(
        session=session,
        violation_type=violation_type,
        severity_score=severity,
        image_snapshot=image_snapshot,
        remarks=remarks
    )

    session.browser_lock_score += severity

    if session.browser_lock_score > 50:
        session.final_verdict = 'High Risk'
    elif session.browser_lock_score > 20:
        session.final_verdict = 'Suspicious'
    else:
        session.final_verdict = 'Safe'

    session.save()

    return Response({
        "status": "success",
        "violation_id": violation.id,
        "browser_lock_score": session.browser_lock_score,
        "final_verdict": session.final_verdict
    })


@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def create_exam_question_api(request):
    """
    Create custom questions manually inside Question Bank
    """
    data = request.data
    subject = data.get('subject', 'PYTHON').upper()
    topic = data.get('topic', '').strip()
    difficulty = data.get('difficulty', 'medium')
    question_text = data.get('question_text', '').strip()
    question_type = data.get('question_type', 'mcq')
    marks = int(data.get('marks', 2))
    choices = data.get('choices', [])

    if not question_text:
        return Response({"error": "question_text is required"}, status=400)

    question = ExamQuestion.objects.create(
        subject=subject,
        topic=topic,
        difficulty=difficulty,
        question_text=question_text,
        question_type=question_type,
        marks=marks
    )

    for choice_data in choices:
        ExamQuestionChoice.objects.create(
            question=question,
            choice_text=choice_data.get('choice_text'),
            is_correct=choice_data.get('is_correct', False)
        )

    return Response({"status": "success", "question_id": question.id})


@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def auto_generate_exam_paper(request):
    """
    Dynamically select random questions matching customized rules
    """
    data = request.data
    title = data.get('title', 'Automated Exam')
    subject = data.get('subject', 'PYTHON').upper()
    easy_count = int(data.get('easy_count', 10))
    medium_count = int(data.get('medium_count', 10))
    hard_count = int(data.get('hard_count', 5))
    duration = int(data.get('duration', 60))

    easy_q = list(ExamQuestion.objects.filter(subject=subject, difficulty='easy'))
    medium_q = list(ExamQuestion.objects.filter(subject=subject, difficulty='medium'))
    hard_q = list(ExamQuestion.objects.filter(subject=subject, difficulty='hard'))

    selected_questions = []
    selected_questions.extend(random.sample(easy_q, min(easy_count, len(easy_q))))
    selected_questions.extend(random.sample(medium_q, min(medium_count, len(medium_q))))
    selected_questions.extend(random.sample(hard_q, min(hard_count, len(hard_q))))

    if not selected_questions:
        return Response({"error": "No questions exist matching criteria in Question Bank"}, status=400)

    total_marks = sum(q.marks for q in selected_questions)

    paper = ExamPaper.objects.create(
        title=title,
        subject=subject,
        duration=duration,
        total_marks=total_marks
    )

    for idx, q in enumerate(selected_questions):
        ExamPaperQuestionRelation.objects.create(
            paper=paper,
            question=q,
            order=idx + 1
        )

    return Response({
        "status": "success",
        "paper_id": paper.id,
        "total_questions": len(selected_questions),
        "total_marks": total_marks
    })


@api_view(['GET'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def get_exam_paper_questions(request, paper_id):
    """
    Fetch all questions and choices belonging to an Exam Paper
    """
    paper = get_object_or_404(ExamPaper, id=paper_id)
    relations = ExamPaperQuestionRelation.objects.filter(paper=paper).order_by('order')

    questions_list = []
    for rel in relations:
        q = rel.question
        choices = []
        correct_index = 0
        for i, c in enumerate(q.choices.all()):
            choices.append({"id": c.id, "choice_text": c.choice_text})
            if c.is_correct:
                correct_index = i
        questions_list.append({
            "id": q.id,
            "question_text": q.question_text,
            "question_type": q.question_type,
            "difficulty": q.difficulty,
            "marks": q.marks,
            "choices": choices,
            "correct": correct_index
        })

    return Response({
        "paper_id": paper.id,
        "title": paper.title,
        "subject": paper.subject,
        "duration": paper.duration,
        "total_marks": paper.total_marks,
        "instructions": paper.instructions,
        "questions": questions_list
    })


# ---------------- AUTOMATED EXAM CONFIG ----------------
@api_view(['GET', 'POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def automated_exam_config_view(request):
    if request.method == 'POST':
        data = request.data
        course_name = str(data.get('course_name', '')).strip()

        if not course_name:
            return Response({"status": "skipped", "message": "course_name required for save"}, status=200)

        course_name_normalized = course_name.upper()

        config = AutomatedExamConfig.objects.filter(course_name__iexact=course_name_normalized).first()

        def safe_int(val, default):
            try:
                if val is None or str(val).strip() == "": return default
                return int(val)
            except: return default

        defaults = {
            'course_name': course_name_normalized,
            'exam_name': data.get('exam_name', 'Daily Assessment'),
            'subjects': data.get('subjects', []),
            'duration': safe_int(data.get('duration'), 80),
            'passing_strategy': data.get('passing_strategy', 'percentage'),
            'requirement': safe_int(data.get('requirement'), 50),
            'question_count': safe_int(data.get('question_count'), 25),
            'marks_per_question': safe_int(data.get('marks_per_question'), 2),
        }

        if config:
            for key, value in defaults.items():
                setattr(config, key, value)
            config.save()
            msg = f"Successfully updated automated config for {course_name_normalized}"
        else:
            config = AutomatedExamConfig.objects.create(**defaults)
            msg = f"Successfully created automated config for {course_name_normalized}"

        return Response({
            "status": "success",
            "config_id": config.id,
            "message": msg
        })

    course_name = request.query_params.get('course_name', '').strip()
    if not course_name:
        return Response({"status": "not_found", "message": "No course_name provided"}, status=200)

    config = AutomatedExamConfig.objects.filter(course_name__iexact=course_name).first()
    if not config:
        return Response({"status": "not_found", "message": "No specific faculty override found for this course."}, status=200)

    return Response({
        "status": "success",
        "exam_name": config.exam_name,
        "subjects": config.subjects,
        "duration": config.duration,
        "passing_strategy": config.passing_strategy,
        "requirement": config.requirement,
        "question_count": config.question_count,
        "marks_per_question": config.marks_per_question,
    })


# ---------------- START EXAM SESSION ----------------
@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([ExamRateThrottle])
def start_exam_session(request):
    data = request.data

    session = ExamSession.objects.create(
        student_name=data.get("student_name"),
        student_email=data.get("student_email"),
        start_time=timezone.now()
    )

    return Response({
        "session_id": session.id
    }, status=status.HTTP_201_CREATED)


# ---------------- SUBMIT ANSWER ----------------
@api_view(['POST'])
@throttle_classes([ExamRateThrottle])
def submit_answer(request, session_id=None):
    data = request.data

    s_id = session_id or data.get("session_id")
    session = get_object_or_404(ExamSession, id=s_id)
    question = get_object_or_404(PythonQuestion, id=data.get("question_id"))

    answer = ExamAnswer.objects.create(
        session=session,
        question=question,
        selected_choice_id=data.get("selected_choice_id"),
        answer_text=data.get("answer_text")
    )

    return Response({
        "answer_id": answer.id
    }, status=status.HTTP_201_CREATED)


# ---------------- END EXAM SESSION ----------------
@api_view(['POST'])
@throttle_classes([ExamRateThrottle])
def end_exam_session(request, session_id):
    session = get_object_or_404(ExamSession, id=session_id)

    session.end_time = timezone.now()
    session.status = "completed"

    answers = ExamAnswer.objects.filter(session=session)

    total_score = 0
    total_marks = 0

    for answer in answers:
        question = answer.question
        total_marks += question.marks

        if answer.selected_choice_id:
            try:
                choice = Choice.objects.get(id=answer.selected_choice_id)
                if choice.is_correct:
                    total_score += question.marks
            except Choice.DoesNotExist:
                pass

    session.score = total_score
    session.total_marks = total_marks
    session.save()

    if session.student_email:
        try:
            threading.Thread(
                target=send_exam_confirmation_email,
                args=(session.student_email, "Placement Portal Exam", session.score, session.total_marks)
            ).start()
        except Exception as e:
            print(f"Error starting email thread: {e}")

    return Response({
        "score": session.score,
        "total_marks": session.total_marks
    })


# ---------------- GET ALL EXAM SESSIONS ----------------
@api_view(['GET'])
@throttle_classes([AuthenticatedUserThrottle])
def get_exam_sessions(request):

    sessions = ExamSession.objects.all().order_by('-created_at')

    data = []

    for session in sessions:
        data.append({
            "id": session.id,
            "student_name": session.student_name,
            "student_email": session.student_email,
            "start_time": session.start_time,
            "end_time": session.end_time,
            "status": session.status,
            "score": session.score,
            "total_marks": session.total_marks,
            "webcam_enabled": session.webcam_enabled,
            "created_at": session.created_at
        })

    return Response(data)


# ---------------- GET ALL QUESTIONS ----------------
@api_view(['GET'])
@throttle_classes([AuthenticatedUserThrottle])
def get_questions(request):

    questions = PythonQuestion.objects.all()

    serializer = PythonQuestionSerializer(questions, many=True)

    return Response(serializer.data)


# ---------------- CREATE QUESTION ----------------
@api_view(['POST'])
@throttle_classes([AuthenticatedUserThrottle])
def create_question(request):

    serializer = PythonQuestionSerializer(data=request.data)

    if serializer.is_valid():

        serializer.save()

        return Response(serializer.data, status=status.HTTP_201_CREATED)

    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


# ---------------- DELETE EXAM SESSION ----------------
@api_view(['DELETE'])
@throttle_classes([AuthenticatedUserThrottle])
def delete_exam_session(request, pk):

    session = get_object_or_404(ExamSession, id=pk)
    session.delete()

    return Response({"message": "Deleted successfully"}, status=status.HTTP_200_OK)


# ---------------- IMPORT EXAM QUESTIONS FILE ----------------

@api_view(['POST'])
@permission_classes([AllowAny])
@authentication_classes([])
@throttle_classes([AuthenticatedUserThrottle])
def import_exam_questions_file(request):
    """
    Parse CSV, Excel, DOCX, or PDF uploaded by faculty,
    and return JSON list of structured questions for the draft.
    """
    if 'file' not in request.FILES:
        return Response({"error": "No file uploaded"}, status=400)

    uploaded_file = request.FILES['file']
    filename = uploaded_file.name.lower()

    questions = []

    try:
        file_bytes = uploaded_file.read()

        if filename.endswith('.csv'):
            file_data = file_bytes.decode('utf-8-sig', errors='ignore')
            reader = csv.reader(io.StringIO(file_data))
            header = next(reader, None)  # skip header
            for row in reader:
                if not row or len(row) < 6:
                    continue
                question_text = str(row[0] or '').strip()
                if not question_text:
                    continue
                options = [str(row[1] or '').strip(), str(row[2] or '').strip(),
                           str(row[3] or '').strip(), str(row[4] or '').strip()]
                correct_indicator = str(row[5] or '').strip()

                correct_idx = 0
                if correct_indicator.isdigit():
                    val = int(correct_indicator)
                    if 1 <= val <= 4:
                        correct_idx = val - 1
                    elif 0 <= val <= 3:
                        correct_idx = val
                    else:
                        correct_idx = 0
                elif correct_indicator.upper() in ['A', 'B', 'C', 'D']:
                    correct_idx = ord(correct_indicator.upper()) - ord('A')
                else:
                    for idx, opt in enumerate(options):
                        if opt and opt.lower() == correct_indicator.lower():
                            correct_idx = idx
                            break

                correct_idx = max(0, min(correct_idx, len([o for o in options if o]) - 1))

                difficulty = str(row[6] or '').strip().lower() if len(row) > 6 else 'medium'
                if difficulty not in ['easy', 'medium', 'hard']:
                    difficulty = 'medium'

                marks = int(str(row[7] or '').strip()) if len(row) > 7 and str(row[7] or '').strip().isdigit() else 1

                questions.append({
                    "question": question_text,
                    "options": options,
                    "correct": correct_idx,
                    "difficulty": difficulty,
                    "marks": marks
                })

        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
            rows = iter(sheet.rows)
            header = next(rows, None)  # skip header
            for r in rows:
                row = [cell.value for cell in r]
                if not row or len(row) < 6 or not row[0]:
                    continue
                question_text = str(row[0] or '').strip()
                if not question_text:
                    continue
                options = [str(row[1] or '').strip(), str(row[2] or '').strip(),
                           str(row[3] or '').strip(), str(row[4] or '').strip()]
                correct_indicator = str(row[5] or '').strip()

                correct_idx = 0
                if correct_indicator.isdigit():
                    val = int(correct_indicator)
                    if 1 <= val <= 4:
                        correct_idx = val - 1
                    elif 0 <= val <= 3:
                        correct_idx = val
                    else:
                        correct_idx = 0
                elif correct_indicator.upper() in ['A', 'B', 'C', 'D']:
                    correct_idx = ord(correct_indicator.upper()) - ord('A')
                else:
                    for idx, opt in enumerate(options):
                        if opt and opt.lower() == correct_indicator.lower():
                            correct_idx = idx
                            break

                correct_idx = max(0, min(correct_idx, len([o for o in options if o]) - 1))

                difficulty = str(row[6] or '').strip().lower() if len(row) > 6 and row[6] else 'medium'
                if difficulty not in ['easy', 'medium', 'hard']:
                    difficulty = 'medium'

                raw_marks = row[7] if len(row) > 7 else None
                marks = int(raw_marks) if raw_marks is not None and str(raw_marks).strip().isdigit() else 1

                questions.append({
                    "question": question_text,
                    "options": options,
                    "correct": correct_idx,
                    "difficulty": difficulty,
                    "marks": marks
                })

        elif filename.endswith('.docx'):
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = "\n".join([p.text for p in doc.paragraphs])
            # Save for inspection
            with open('C:/SSSIT MAIN PROJECT/python12to1pmmainproject/parsed_text_debug.txt', 'w', encoding='utf-8') as f:
                f.write(full_text)
            questions = parse_raw_text_to_questions(full_text)

        elif filename.endswith('.pdf'):
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            except ImportError:
                import PyPDF2 as pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            full_text = ""
            for page in reader.pages:
                try:
                    text = page.extract_text()
                except Exception:
                    text = ""
                if text:
                    full_text += text + "\n"
            # Collapse multiple spaces (common PDF artifact)
            full_text = re.sub(r'[ \t]+', ' ', full_text)
            # Save for inspection
            with open('C:/SSSIT MAIN PROJECT/python12to1pmmainproject/parsed_text_debug.txt', 'w', encoding='utf-8') as f:
                f.write(full_text)
            questions = parse_raw_text_to_questions(full_text)

        else:
            return Response({"error": "Unsupported file format. Please upload CSV, Excel, Word (docx), or PDF."}, status=400)

    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({"error": f"Failed to parse file: {str(e)}"}, status=500)

    return Response({"status": "success", "questions": questions})

def parse_raw_text_to_questions(text):
    """
    Parse structured Q&A text from PDF/DOCX into question dicts.

    Supported Answer Formats:
    1. "Answer: B", "Ans: C", "Correct Answer: D", "Correct: A"
    2. Starred/Bracketed Options: "*A) text", "B)* text", "[C] text"
    3. End-of-document key tables: "1-B, 2-D, 3-A" or "1. B, 2. D"
    4. Compact no-space options: "A)2752 B)2746 C)2734 D)2718"
    """
    import logging
    logger = logging.getLogger(__name__)

    questions = []

    # ────────────────────────────────────────────────────────────────────────
    # STEP 1: Find global answer key table (end-of-doc style: 1-B, 2-C ...)
    # ────────────────────────────────────────────────────────────────────────
    global_answers = {}
    ans_key_re = re.compile(r'(?<!\d)(\d{1,3})\s*[-\u2013\u2014\.:\)]\s*([A-Da-d])(?!\w)')
    all_pairs = list(ans_key_re.finditer(text))

    if len(all_pairs) >= 3:
        clusters = []
        cur = [all_pairs[0]]
        for j in range(1, len(all_pairs)):
            if all_pairs[j].start() - all_pairs[j-1].start() < 300:
                cur.append(all_pairs[j])
            else:
                clusters.append(cur)
                cur = [all_pairs[j]]
        clusters.append(cur)

        largest = max(clusters, key=len)
        if len(largest) >= 3:
            for pair in largest:
                global_answers[int(pair.group(1))] = ord(pair.group(2).upper()) - ord('A')
            logger.info(f"Global answer key found: {global_answers}")

    # ────────────────────────────────────────────────────────────────────────
    # STEP 2: Split text into per-question blocks
    # ────────────────────────────────────────────────────────────────────────
    q_re = re.compile(
        r'(?:^|\n)\s*(?:Q(?:uestion)?\s*\.?\s*(\d+)\.?|(\d+)\s*[\)\.\-:])\s*',
        re.IGNORECASE
    )
    matches = list(q_re.finditer(text))
    if not matches:
        logger.warning("No question markers found in document.")
        return questions

    logger.info(f"Found {len(matches)} question blocks")

    for i, m in enumerate(matches):
        q_num_str = m.group(1) or m.group(2)
        q_num = int(q_num_str) if (q_num_str and q_num_str.isdigit()) else (i + 1)

        block_start = m.end()
        block_end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        block = text[block_start:block_end].strip()
        if not block:
            continue

        # ──────────────────────────────────────────────────────────────────
        # STEP 3: Extract inline answer indicator BEFORE parsing options
        # Matches: "Answer: B", "Ans.B", "Answer=D", "Correct: C", "Key: A"
        # ──────────────────────────────────────────────────────────────────
        inline_ans_re = re.compile(
            r'(?i)\b(?:ans(?:wer)?(?:\s*key)?|correct(?:\s*answer)?|key)\s*[=:\.\-\u2013\u2014]?\s*([A-Da-d])\b'
        )
        inline_m = inline_ans_re.search(block)
        found_letter = None
        clean_block = block

        if inline_m:
            found_letter = inline_m.group(1).upper()
            clean_block = (block[:inline_m.start()] + block[inline_m.end():]).strip()
            logger.debug(f"Q{q_num}: Inline answer '{found_letter}'")

        # ──────────────────────────────────────────────────────────────────
        # STEP 4: Parse options (handles both spaced and compact formats)
        # ──────────────────────────────────────────────────────────────────
        opt_marker_re = re.compile(r'([A-Da-d])\s*[\)\.\:]', re.MULTILINE)
        raw_markers = list(opt_marker_re.finditer(clean_block))

        valid_opts = []
        expected = ['A', 'B', 'C', 'D']
        ei = 0
        for om in raw_markers:
            lbl = om.group(1).upper()
            if ei < len(expected) and lbl == expected[ei]:
                valid_opts.append(om)
                ei += 1
                if ei == 4:
                    break

        options = []
        question_body = clean_block

        if len(valid_opts) >= 2:
            question_body = clean_block[:valid_opts[0].start()].strip()
            for k, om in enumerate(valid_opts):
                s = om.end()
                e = valid_opts[k + 1].start() if k + 1 < len(valid_opts) else len(clean_block)
                opt_text = clean_block[s:e].strip()
                opt_text = re.sub(r'\s+', ' ', opt_text).strip()
                opt_text = re.sub(
                    r'(?i)\s*\b(?:ans(?:wer)?|correct|key)\s*[=:\.\-\u2013\u2014]?\s*[A-Da-d]\b.*$',
                    '', opt_text
                ).strip()
                options.append(opt_text)
        else:
            opt_line_re = re.compile(r'^([A-Da-d])\s*[\)\.\:]\s*(.+)$')
            q_lines = []
            for line in clean_block.split('\n'):
                line = line.strip()
                if not line:
                    continue
                lm = opt_line_re.match(line)
                if lm:
                    options.append(lm.group(2).strip())
                else:
                    q_lines.append(line)
            question_body = ' '.join(q_lines).strip()

        if not question_body:
            continue

        # ──────────────────────────────────────────────────────────────────
        # STEP 5: Determine correct answer index
        # ──────────────────────────────────────────────────────────────────
        correct_idx = None

        if found_letter is not None:
            correct_idx = ord(found_letter) - ord('A')

        if correct_idx is None:
            star_re = re.compile(
                r'(?i)[\*\[\u2713\u2714]\s*([A-Da-d])\s*[\)\]\*]'
                r'|([A-Da-d])\s*[\)\.]?\s*[\*\[\u2713\u2714]'
            )
            sm = star_re.search(block)
            if sm:
                lbl = (sm.group(1) or sm.group(2) or '').upper()
                if lbl in 'ABCD':
                    correct_idx = ord(lbl) - ord('A')
                    logger.debug(f"Q{q_num}: Star/bracket answer '{lbl}'")

        if correct_idx is None and q_num in global_answers:
            correct_idx = global_answers[q_num]
            logger.debug(f"Q{q_num}: Global key -> {chr(ord('A') + correct_idx)}")

        if correct_idx is None:
            full_re = re.compile(
                r'(?i)\b(?:ans(?:wer)?|correct)\s*[=:\.\-\u2013\u2014]?\s*(.+)$',
                re.MULTILINE
            )
            fm = full_re.search(block)
            if fm:
                ans_text = fm.group(1).strip().rstrip('.')
                for k, opt in enumerate(options):
                    if opt and ans_text.lower() in opt.lower():
                        correct_idx = k
                        logger.debug(f"Q{q_num}: Answer text matched option {k}")
                        break

        if correct_idx is None:
            logger.warning(f"Q{q_num}: No answer detected, defaulting to A. Block: {block[:100]!r}")
            correct_idx = 0

        if options:
            correct_idx = max(0, min(correct_idx, len(options) - 1))

        # ──────────────────────────────────────────────────────────────────
        # STEP 6: Save question if valid
        # ──────────────────────────────────────────────────────────────────
        if len(options) >= 2:
            padded = (options + [''] * 4)[:4]
            questions.append({
                'question': question_body,
                'options': padded,
                'correct': correct_idx,
                'difficulty': 'medium',
                'marks': 1
            })
            logger.debug(f"Q{q_num}: Saved. Ans={chr(ord('A') + correct_idx)}, opts={options}")
        else:
            logger.warning(f"Q{q_num}: Skipped — fewer than 2 options. Block: {block[:80]!r}")

    logger.info(f"Total questions parsed: {len(questions)}")
    return questions

